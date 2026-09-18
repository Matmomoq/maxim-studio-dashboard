"""Flask application serving the amoCRM marketing dashboard."""

from __future__ import annotations

import hmac
import io
import os
import secrets
from calendar import monthrange
from datetime import date, datetime, timedelta
from decimal import Decimal
from functools import wraps
from pathlib import Path
from typing import Any, Callable, Dict, Optional, Tuple
from urllib.parse import urlsplit

from flask import (
    Flask,
    g,
    Response,
    abort,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    stream_with_context,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from werkzeug.exceptions import HTTPException

from .login_security import LoginLimiter
from amocrm_client import load_env
from .analytics import (
    MIN_DASHBOARD_DATE,
    MOSCOW,
    DashboardAnalytics,
    details_to_csv,
    marketing_reconciliation_to_csv,
    verification_to_csv,
)
from .marketing_expenses import (
    DIRECTIONS,
    MarketingExpenseRepository,
    MarketingSchemaMissing,
    PeriodExpenseInput,
    calculate_period_preview,
    validate_period_input,
)
from .marketing_analytics import calculate_marketing_performance
from .trend_analytics import calculate_trend_data
from .call_analytics import CallAnalytics
from .call_recordings_demo import CallRecordingDemo, RecordingUnavailable
from .funnel_analytics import FunnelAnalytics
from .yclients_analytics import (
    MIN_YCLIENTS_DATE,
    YclientsAnalytics,
    color_reconciliation_to_csv,
)
from .budget_planning import (
    BUDGET_DIRECTIONS,
    BUDGET_SOURCES,
    BudgetSchemaMissing,
    MarketingBudgetRepository,
    calculate_budget_rows,
)
from .plan_fact import (
    DIRECTIONS as PLAN_FACT_DIRECTIONS,
    PLAN_FACT_START,
    PlanFactAnalytics,
    PlanFactRepository,
    PlanFactSchemaMissing,
    metric_catalog,
    parse_plan_month,
)


def _moscow_today() -> date:
    return datetime.now(MOSCOW).date()


def _parse_period() -> Tuple[date, date]:
    today = _moscow_today()
    default_start = today.replace(day=1)
    try:
        start = date.fromisoformat(request.args.get("date_from", ""))
    except ValueError:
        start = max(default_start, MIN_DASHBOARD_DATE)
    try:
        end = date.fromisoformat(request.args.get("date_to", ""))
    except ValueError:
        end = today
    if start < MIN_DASHBOARD_DATE:
        start = MIN_DASHBOARD_DATE
    if end < start:
        raise ValueError("Дата окончания не может быть раньше даты начала")
    if (end - start).days > 730:
        raise ValueError("Выберите период не более двух лет")
    return start, end


def _parse_yclients_period() -> Tuple[date, date]:
    today = _moscow_today()
    default_start = today.replace(day=1)
    try:
        start = date.fromisoformat(request.args.get("date_from", ""))
    except ValueError:
        start = default_start
    try:
        end = date.fromisoformat(request.args.get("date_to", ""))
    except ValueError:
        end = today
    if start < MIN_YCLIENTS_DATE:
        start = MIN_YCLIENTS_DATE
    if end < start:
        raise ValueError("Дата окончания не может быть раньше даты начала")
    if (end - start).days > 730:
        raise ValueError("Выберите период не более двух лет")
    return start, end


def _filters(
    keys=("branch", "direction", "offer", "source"),
) -> Dict[str, list[str]]:
    return {
        key: request.args.getlist(key) or ["Все"]
        for key in keys
    }


def _optional_date(name: str) -> Optional[date]:
    raw = request.args.get(name, "").strip()
    if not raw:
        return None
    try:
        value = date.fromisoformat(raw)
    except ValueError as error:
        raise ValueError("Некорректная дата") from error
    if value < MIN_DASHBOARD_DATE:
        return MIN_DASHBOARD_DATE
    return value


def _parse_month(raw: str) -> Tuple[date, date]:
    try:
        year_text, month_text = str(raw or "").split("-", 1)
        year = int(year_text)
        month = int(month_text)
        start = date(year, month, 1)
    except (TypeError, ValueError) as error:
        raise ValueError("Выберите корректный месяц расходов") from error
    if start < MIN_DASHBOARD_DATE:
        raise ValueError("Расходы доступны начиная с июля 2026 года")
    return start, date(year, month, monthrange(year, month)[1])


def _month_label(value: date) -> str:
    names = (
        "январь", "февраль", "март", "апрель", "май", "июнь",
        "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь",
    )
    return f"{names[value.month - 1].capitalize()} {value.year}"


def _xlsx_bytes(details, start: date, end: date) -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Эффективность"
    headers = [
        "Филиал",
        "Направление",
        "Оффер",
        "Источник",
        "Период с",
        "Период по",
        "Лиды",
        "Записи",
        "Конверсия",
    ]
    sheet.append(headers)
    for row in details:
        sheet.append(
            [
                row["branch"],
                row["direction"],
                row["offer"],
                row["source"],
                start,
                end,
                row["leads"],
                row["bookings"],
                row["conversion"] / 100,
            ]
        )
    header_fill = PatternFill("solid", fgColor="172A26")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [22, 18, 26, 18, 14, 14, 12, 12, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for cell in sheet["I"][1:]:
        cell.number_format = "0.00%"
    for column in ("E", "F"):
        for cell in sheet[column][1:]:
            cell.number_format = "DD.MM.YYYY"
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def create_app(env_path: str = ".env") -> Flask:
    app = Flask(
        __name__,
        template_folder="templates",
        static_folder="static",
    )
    app.config["JSON_AS_ASCII"] = False
    asset_directory = Path(app.static_folder)
    inline_css = (asset_directory / "styles.css").read_text(encoding="utf-8")
    inline_js = (asset_directory / "app.js").read_text(encoding="utf-8")
    inline_verification_js = (asset_directory / "verification.js").read_text(
        encoding="utf-8"
    )
    inline_expenses_js = (asset_directory / "expenses.js").read_text(
        encoding="utf-8"
    )
    inline_budgets_js = (asset_directory / "budgets.js").read_text(
        encoding="utf-8"
    )
    inline_calls_js = (asset_directory / "calls.js").read_text(encoding="utf-8")
    inline_call_recordings_demo_js = (
        asset_directory / "call_recordings_demo.js"
    ).read_text(encoding="utf-8")
    inline_funnel_js = (asset_directory / "funnel.js").read_text(encoding="utf-8")
    inline_yclients_js = (asset_directory / "yclients.js").read_text(encoding="utf-8")
    inline_plans_js = (asset_directory / "plans.js").read_text(encoding="utf-8")
    inline_plan_fact_js = (asset_directory / "plan_fact.js").read_text(encoding="utf-8")
    analytics = DashboardAnalytics(env_path)
    app.extensions["dashboard_analytics"] = analytics
    app.extensions["marketing_expenses"] = MarketingExpenseRepository(env_path)
    app.extensions["marketing_budgets"] = MarketingBudgetRepository(env_path)
    app.extensions["call_analytics"] = CallAnalytics(env_path)
    app.extensions["call_recordings_demo"] = CallRecordingDemo(env_path)
    app.extensions["funnel_analytics"] = FunnelAnalytics(env_path)
    app.extensions["yclients_analytics"] = YclientsAnalytics(env_path)
    app.extensions["plan_fact_repository"] = PlanFactRepository(env_path)
    app.extensions["plan_fact_analytics"] = PlanFactAnalytics(
        env_path, repository=app.extensions["plan_fact_repository"]
    )
    env = load_env(env_path)
    # Process-level dashboard settings are useful for local previews and
    # hosting panels, while the database/API settings continue to come from
    # the selected .env file.
    env.update(
        {
            key: value
            for key, value in os.environ.items()
            if key.startswith("DASHBOARD_")
        }
    )
    auth_required = env.get("DASHBOARD_AUTH_REQUIRED", "true").lower() not in {
        "0",
        "false",
        "no",
    }
    expected_user = env.get("DASHBOARD_USERNAME", "")
    expected_password = env.get("DASHBOARD_PASSWORD", "")
    credentials_configured = bool(expected_user and expected_password)
    secret_key = env.get("DASHBOARD_SECRET_KEY", "")
    if auth_required and credentials_configured and not secret_key:
        raise RuntimeError("DASHBOARD_SECRET_KEY must be configured independently")
    limiter = LoginLimiter(
        env.get("DASHBOARD_LOGIN_LIMIT_DB", str(Path(env_path).resolve().parent / ".dashboard-login.sqlite3")),
        secret_key or "local-preview",
    )
    app.extensions["login_limiter"] = limiter
    app.config.update(
        SECRET_KEY=secret_key or secrets.token_hex(32),
        PERMANENT_SESSION_LIFETIME=timedelta(
            hours=max(1, int(env.get("DASHBOARD_SESSION_HOURS", "12")))
        ),
        SESSION_COOKIE_NAME="amo_dashboard_session",
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=env.get(
            "DASHBOARD_COOKIE_SECURE", "true"
        ).lower() not in {"0", "false", "no"},
    )

    def safe_next(value: str) -> str:
        candidate = value or "/"
        parsed = urlsplit(candidate)
        if (
            not candidate.startswith("/")
            or candidate.startswith("//")
            or parsed.scheme
            or parsed.netloc
        ):
            return "/"
        return candidate

    def csrf_token(name: str) -> str:
        token = session.get(name)
        if not token:
            token = secrets.token_urlsafe(32)
            session[name] = token
        return token

    def require_csrf(name: str) -> None:
        submitted = request.headers.get("X-CSRF-Token", "")
        stored = session.get(name, "")
        if not (
            submitted
            and stored
            and hmac.compare_digest(submitted, stored)
        ):
            abort(400)

    def json_payload() -> Dict[str, Any]:
        payload = request.get_json(silent=True)
        if not isinstance(payload, dict):
            raise ValueError("Не удалось прочитать данные формы")
        return payload

    def expense_branches() -> list[str]:
        return sorted(analytics.classifier.mapping.get("branches", {}).keys())

    def expense_repository() -> MarketingExpenseRepository:
        return app.extensions["marketing_expenses"]

    def expense_preview_records(start: date, end: date):
        """Expense matching uses campaign tags, not the visit branch field."""
        return analytics.marketing_tag_records(start, end)

    def budget_repository() -> MarketingBudgetRepository:
        return app.extensions["marketing_budgets"]

    def calls_analytics() -> CallAnalytics:
        return app.extensions["call_analytics"]

    def call_recordings_demo() -> CallRecordingDemo:
        return app.extensions["call_recordings_demo"]

    def funnel_analytics() -> FunnelAnalytics:
        return app.extensions["funnel_analytics"]

    def yclients_analytics() -> YclientsAnalytics:
        return app.extensions["yclients_analytics"]

    def plan_fact_repository() -> PlanFactRepository:
        return app.extensions["plan_fact_repository"]

    def plan_fact_analytics() -> PlanFactAnalytics:
        return app.extensions["plan_fact_analytics"]

    def is_authenticated() -> bool:
        return bool(
            credentials_configured
            and session.get("dashboard_authenticated")
            and hmac.compare_digest(
                str(session.get("dashboard_user", "")), expected_user
            )
        )

    def requires_auth(handler: Callable[..., Any]):
        @wraps(handler)
        def wrapped(*args, **kwargs):
            if not auth_required:
                return handler(*args, **kwargs)
            if not credentials_configured:
                if request.path.startswith("/api/"):
                    return jsonify({"error": "Авторизация не настроена"}), 503
                return redirect(url_for("login"))
            if not is_authenticated():
                if request.path.startswith("/api/"):
                    return jsonify({"error": "Требуется авторизация"}), 401
                target = request.full_path.rstrip("?")
                return redirect(url_for("login", next=target))
            return handler(*args, **kwargs)

        return wrapped

    @app.before_request
    def prepare_csp_nonce():
        g.csp_nonce = secrets.token_urlsafe(24)

    @app.context_processor
    def security_template_context():
        return {"csp_nonce": g.csp_nonce}

    @app.after_request
    def security_headers(response):
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Referrer-Policy"] = "same-origin"
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; style-src 'self' 'unsafe-inline'; "
            f"script-src 'self' 'nonce-{g.csp_nonce}'; object-src 'none'; connect-src 'self'; "
            "img-src 'self' data:; media-src 'self'; frame-ancestors 'none'; "
            "base-uri 'self'; form-action 'self'"
        )
        if request.path != "/health":
            response.headers["Cache-Control"] = "no-store"
        return response

    @app.errorhandler(ValueError)
    def value_error(error):
        return jsonify({"error": str(error)}), 400

    @app.errorhandler(MarketingSchemaMissing)
    def marketing_schema_missing(error):
        return jsonify({"error": str(error)}), 503

    @app.errorhandler(BudgetSchemaMissing)
    def budget_schema_missing(error):
        return jsonify({"error": str(error)}), 503

    @app.errorhandler(PlanFactSchemaMissing)
    def plan_fact_schema_missing(error):
        return jsonify({"error": str(error)}), 503

    @app.errorhandler(RecordingUnavailable)
    def recording_unavailable(error):
        return jsonify({"error": str(error)}), 502

    @app.errorhandler(Exception)
    def unexpected_error(error):
        if isinstance(error, HTTPException):
            return error
        app.logger.exception("Dashboard request failed")
        return jsonify({"error": "Не удалось получить данные из базы"}), 500

    @app.route("/login", methods=["GET", "POST"])
    def login():
        next_url = safe_next(request.values.get("next", "/"))
        if not auth_required or is_authenticated():
            return redirect(next_url)
        error = ""
        status = 200
        if not credentials_configured:
            error = (
                "Вход ещё не настроен. Заполните DASHBOARD_USERNAME и "
                "DASHBOARD_PASSWORD в файле .env и перезапустите приложение."
            )
            status = 503
        elif request.method == "POST":
            submitted_token = request.form.get("csrf_token", "")
            stored_token = session.get("login_csrf", "")
            if not (
                submitted_token
                and stored_token
                and hmac.compare_digest(submitted_token, stored_token)
            ):
                abort(400)
            username = request.form.get("username", "")
            password = request.form.get("password", "")
            # Count attempts atomically before password verification across Passenger workers.
            retry_after = limiter.attempt(request.remote_addr or "unknown", username)
            if retry_after:
                response = app.make_response((render_template(
                    "login.html", inline_css=inline_css,
                    csrf_token=csrf_token("login_csrf"),
                    error="Слишком много попыток входа. Попробуйте через несколько минут.",
                    configured=credentials_configured, next_url=next_url,
                ), 429))
                response.headers["Retry-After"] = str(retry_after)
                return response
            valid = hmac.compare_digest(username.encode(), expected_user.encode()) and hmac.compare_digest(
                password.encode(), expected_password.encode()
            )
            if valid:
                session.clear()
                session["dashboard_authenticated"] = True
                session["dashboard_user"] = expected_user
                session.permanent = True
                return redirect(next_url)
            app.logger.warning("Dashboard login rejected")
            error = "Неверный логин или пароль"
            status = 401
        return (
            render_template(
                "login.html",
                inline_css=inline_css,
                csrf_token=csrf_token("login_csrf"),
                error=error,
                configured=credentials_configured,
                next_url=next_url,
            ),
            status,
        )

    @app.post("/logout")
    @requires_auth
    def logout():
        submitted_token = request.form.get("csrf_token", "")
        stored_token = session.get("logout_csrf", "")
        if not (
            submitted_token
            and stored_token
            and hmac.compare_digest(submitted_token, stored_token)
        ):
            abort(400)
        session.clear()
        return redirect(url_for("login"))

    @app.get("/")
    @requires_auth
    def index():
        return render_template(
            "index.html",
            min_date=MIN_DASHBOARD_DATE.isoformat(),
            today=_moscow_today().isoformat(),
            inline_css=inline_css,
            inline_js=inline_js,
            logout_csrf=csrf_token("logout_csrf"),
        )

    @app.get("/verification")
    @requires_auth
    def verification_page():
        return render_template(
            "verification.html",
            min_date=MIN_DASHBOARD_DATE.isoformat(),
            today=_moscow_today().isoformat(),
            inline_css=inline_css,
            inline_js=inline_verification_js,
            logout_csrf=csrf_token("logout_csrf"),
        )

    @app.get("/expenses")
    @requires_auth
    def expenses_page():
        return render_template(
            "expenses.html",
            min_date=MIN_DASHBOARD_DATE.isoformat(),
            today=_moscow_today().isoformat(),
            expense_branches=expense_branches(),
            expense_directions=list(DIRECTIONS),
            inline_css=inline_css,
            inline_js=inline_expenses_js,
            expense_csrf=csrf_token("expense_csrf"),
            logout_csrf=csrf_token("logout_csrf"),
        )

    @app.get("/budgets")
    @requires_auth
    def budgets_page():
        return render_template(
            "budgets.html",
            min_date=MIN_DASHBOARD_DATE.isoformat(),
            today=_moscow_today().isoformat(),
            inline_css=inline_css,
            inline_js=inline_budgets_js,
            budget_csrf=csrf_token("budget_csrf"),
            logout_csrf=csrf_token("logout_csrf"),
        )

    @app.get("/plans")
    @requires_auth
    def plans_page():
        return render_template(
            "plans.html",
            min_date=PLAN_FACT_START.isoformat(),
            today=_moscow_today().isoformat(),
            inline_css=inline_css,
            inline_js=inline_plans_js,
            plan_csrf=csrf_token("plan_fact_csrf"),
            logout_csrf=csrf_token("logout_csrf"),
        )

    @app.get("/calls")
    @requires_auth
    def calls_page():
        return render_template(
            "calls.html",
            min_date=MIN_DASHBOARD_DATE.isoformat(),
            today=_moscow_today().isoformat(),
            inline_css=inline_css,
            inline_js=inline_calls_js,
            logout_csrf=csrf_token("logout_csrf"),
        )

    @app.get("/call-recordings-demo")
    @requires_auth
    def call_recordings_demo_page():
        return render_template(
            "call_recordings_demo.html",
            inline_css=inline_css,
            inline_js=inline_call_recordings_demo_js,
            logout_csrf=csrf_token("logout_csrf"),
        )

    @app.get("/funnel")
    @requires_auth
    def funnel_page():
        return render_template(
            "funnel.html",
            min_date=MIN_DASHBOARD_DATE.isoformat(),
            today=_moscow_today().isoformat(),
            inline_css=inline_css,
            inline_js=inline_funnel_js,
            logout_csrf=csrf_token("logout_csrf"),
        )

    @app.get("/yclients")
    @requires_auth
    def yclients_page():
        return render_template(
            "yclients.html",
            min_date=MIN_YCLIENTS_DATE.isoformat(),
            today=_moscow_today().isoformat(),
            inline_css=inline_css,
            inline_js=inline_yclients_js,
            logout_csrf=csrf_token("logout_csrf"),
        )

    @app.get("/yclients/plan-fact")
    @requires_auth
    def yclients_plan_fact_page():
        return render_template(
            "yclients_plan_fact.html",
            min_date=PLAN_FACT_START.isoformat(),
            today=_moscow_today().isoformat(),
            inline_css=inline_css,
            inline_js=inline_plan_fact_js,
            logout_csrf=csrf_token("logout_csrf"),
        )

    @app.get("/health")
    def health():
        return jsonify({"status": "ok"})

    @app.get("/api/filters")
    @requires_auth
    def filters():
        start = MIN_DASHBOARD_DATE
        end = _moscow_today()
        records = analytics.records(start, end)
        return jsonify(
            {
                "options": analytics.filter_options(records),
                "min_date": start.isoformat(),
                "max_date": end.isoformat(),
                "last_sync": analytics.last_sync(),
            }
        )

    @app.get("/api/expenses")
    @requires_auth
    def expenses_api():
        start, end = _parse_month(
            request.args.get("month") or _moscow_today().strftime("%Y-%m")
        )
        schema_ready = True
        try:
            rates = expense_repository().list_rates()
            expenses = expense_repository().list_period_expenses(start, end)
        except MarketingSchemaMissing:
            schema_ready = False
            rates = []
            expenses = []
        if expenses:
            record_start = min(date.fromisoformat(row["period_from"]) for row in expenses)
            record_end = max(date.fromisoformat(row["period_to"]) for row in expenses)
            records = expense_preview_records(record_start, record_end)
            for row in expenses:
                preview_input = PeriodExpenseInput(
                    source=row["source"],
                    branch=row["branch"],
                    direction=row["direction"],
                    period_from=date.fromisoformat(row["period_from"]),
                    period_to=date.fromisoformat(row["period_to"]),
                    amount=Decimal(str(row["amount"])),
                    status=row["status"],
                    comment=row["comment"],
                )
                row["preview"] = calculate_period_preview(records, preview_input)
        return jsonify(
            {
                "month": start.strftime("%Y-%m"),
                "month_from": start.isoformat(),
                "month_to": end.isoformat(),
                "branches": expense_branches(),
                "directions": list(DIRECTIONS),
                "rates": rates,
                "expenses": expenses,
                "schema_ready": schema_ready,
                "last_sync": analytics.last_sync(),
            }
        )

    @app.get("/api/plans")
    @requires_auth
    def plans_api():
        month = parse_plan_month(
            request.args.get("month") or _moscow_today().strftime("%Y-%m")
        )
        schema_ready = True
        try:
            plans = plan_fact_repository().list_plans(month)
        except PlanFactSchemaMissing:
            plans = []
            schema_ready = False
        return jsonify(
            {
                "month": month.strftime("%Y-%m"),
                "plans": plans,
                "metrics": metric_catalog(),
                "options": {
                    "branches": expense_branches(),
                    "directions": list(PLAN_FACT_DIRECTIONS),
                },
                "schema_ready": schema_ready,
            }
        )

    @app.post("/api/plans")
    @requires_auth
    def plan_create_api():
        require_csrf("plan_fact_csrf")
        saved = plan_fact_repository().save_plan(
            json_payload(),
            branches=expense_branches(),
            user=str(session.get("dashboard_user", "dashboard")),
        )
        return jsonify({"plan": saved}), 201

    @app.put("/api/plans/<int:plan_id>")
    @requires_auth
    def plan_update_api(plan_id: int):
        require_csrf("plan_fact_csrf")
        saved = plan_fact_repository().save_plan(
            json_payload(),
            branches=expense_branches(),
            user=str(session.get("dashboard_user", "dashboard")),
            plan_id=plan_id,
        )
        return jsonify({"plan": saved})

    @app.delete("/api/plans/<int:plan_id>")
    @requires_auth
    def plan_delete_api(plan_id: int):
        require_csrf("plan_fact_csrf")
        deleted = plan_fact_repository().delete_plan(
            plan_id, user=str(session.get("dashboard_user", "dashboard"))
        )
        return jsonify({"plan": deleted})

    @app.get("/api/plans/<int:plan_id>/history")
    @requires_auth
    def plan_history_api(plan_id: int):
        return jsonify({"history": plan_fact_repository().history(plan_id)})

    @app.post("/api/expenses/preview")
    @requires_auth
    def expenses_preview_api():
        require_csrf("expense_csrf")
        expense = validate_period_input(
            json_payload(), branches=expense_branches()
        )
        records = expense_preview_records(expense.period_from, expense.period_to)
        return jsonify(calculate_period_preview(records, expense))

    @app.get("/api/expenses/reconciliation")
    @requires_auth
    def expenses_reconciliation_api():
        start, end = _parse_period()
        result = analytics.marketing_reconciliation(
            start,
            end,
            source=request.args.get("source", "").strip(),
            branch=request.args.get("branch", "").strip(),
            direction=request.args.get("direction", "").strip(),
        )
        result["period"] = {"from": start.isoformat(), "to": end.isoformat()}
        result["last_sync"] = analytics.last_sync()
        return jsonify(result)

    @app.get("/api/expenses/reconciliation.csv")
    @requires_auth
    def expenses_reconciliation_export():
        start, end = _parse_period()
        result = analytics.marketing_reconciliation(
            start,
            end,
            source=request.args.get("source", "").strip(),
            branch=request.args.get("branch", "").strip(),
            direction=request.args.get("direction", "").strip(),
        )
        return Response(
            marketing_reconciliation_to_csv(result),
            mimetype="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": (
                    f'attachment; filename="marketing_tags_{start}_{end}.csv"'
                )
            },
        )

    @app.post("/api/expenses/rates")
    @requires_auth
    def expenses_rate_create_api():
        require_csrf("expense_csrf")
        created = expense_repository().add_rate(
            json_payload(), user=str(session.get("dashboard_user", "dashboard"))
        )
        return jsonify({"rate": created}), 201

    @app.post("/api/expenses/periods")
    @requires_auth
    def expenses_period_create_api():
        require_csrf("expense_csrf")
        payload = json_payload()
        created = expense_repository().add_period_expense(
            payload,
            branches=expense_branches(),
            user=str(session.get("dashboard_user", "dashboard")),
        )
        expense = validate_period_input(payload, branches=expense_branches())
        records = expense_preview_records(expense.period_from, expense.period_to)
        created["preview"] = calculate_period_preview(records, expense)
        return jsonify({"expense": created}), 201

    @app.post("/api/expenses/periods/<int:expense_id>/<status>")
    @requires_auth
    def expenses_period_status_api(expense_id: int, status: str):
        require_csrf("expense_csrf")
        updated = expense_repository().set_period_status(
            expense_id,
            status,
            user=str(session.get("dashboard_user", "dashboard")),
        )
        return jsonify({"expense": updated})

    @app.get("/api/budgets")
    @requires_auth
    def budgets_api():
        start, end = _parse_month(
            request.args.get("month") or _moscow_today().strftime("%Y-%m")
        )
        schema_ready = True
        try:
            budgets = budget_repository().list_budgets(start)
        except BudgetSchemaMissing:
            budgets = []
            schema_ready = False
        try:
            rates = expense_repository().list_rates()
            period_expenses = expense_repository().list_period_expenses(start, end)
        except MarketingSchemaMissing:
            rates = []
            period_expenses = []
        result = calculate_budget_rows(
            analytics.records(start, end),
            budgets,
            rates=rates,
            period_expenses=period_expenses,
            month_start=start,
            month_end=end,
            filters={
                key: request.args.getlist(key) or ["Все"]
                for key in ("source", "branch", "direction")
            },
        )
        selected_budget_filters = {
            key: {
                value for value in request.args.getlist(key)
                if value and value != "Все"
            }
            for key in ("source", "branch", "direction")
        }
        budget_history = [
            row for row in budgets
            if all(
                not selected_budget_filters[key]
                or row[key] in selected_budget_filters[key]
                for key in ("source", "branch", "direction")
            )
        ]
        return jsonify(
            {
                **result,
                "budgets": budget_history,
                "month": start.strftime("%Y-%m"),
                "month_label": _month_label(start),
                "options": {
                    "sources": list(BUDGET_SOURCES),
                    "branches": expense_branches(),
                    "directions": list(BUDGET_DIRECTIONS),
                },
                "schema_ready": schema_ready,
                "last_sync": analytics.last_sync(),
            }
        )

    @app.get("/api/calls")
    @requires_auth
    def calls_api():
        start, end = _parse_period()
        result = calls_analytics().report(
            start,
            end,
            team=request.args.get("team", "cc"),
            city=request.args.get("city", "Все"),
            user=request.args.get("user", "Все"),
        )
        result["last_sync"] = analytics.last_sync()
        return jsonify(result)

    @app.get("/api/call-recordings-demo")
    @requires_auth
    def call_recordings_demo_api():
        return jsonify(call_recordings_demo().list_calls())

    @app.get("/api/call-recordings-demo/<int:note_id>/audio")
    @requires_auth
    def call_recording_demo_audio(note_id: int):
        upstream = call_recordings_demo().open_audio(
            note_id, request.headers.get("Range")
        )

        def audio_chunks():
            try:
                yield from upstream.iter_content(chunk_size=64 * 1024)
            finally:
                upstream.close()

        headers = {
            key: upstream.headers[key]
            for key in (
                "Content-Type", "Content-Length", "Content-Range", "Accept-Ranges"
            )
            if upstream.headers.get(key)
        }
        headers.setdefault("Content-Type", "audio/mpeg")
        headers["Content-Disposition"] = f'inline; filename="call-{note_id}.mp3"'
        return Response(
            stream_with_context(audio_chunks()),
            status=upstream.status_code,
            headers=headers,
        )

    @app.get("/api/funnel")
    @requires_auth
    def funnel_api():
        start, end = _parse_period()
        filters = _filters()
        records = analytics.records(start, end)
        include_drafts = request.args.get("include_drafts", "false").lower() in {
            "1", "true", "yes"
        }
        try:
            rates = expense_repository().list_rates()
            expenses = expense_repository().list_period_expenses(start, end)
        except MarketingSchemaMissing:
            rates = []
            expenses = []
        allocation_start = start
        allocation_end = end
        if expenses:
            allocation_start = min(
                allocation_start,
                min(date.fromisoformat(row["period_from"]) for row in expenses),
            )
            allocation_end = max(
                allocation_end,
                max(date.fromisoformat(row["period_to"]) for row in expenses),
            )
        allocation_records = analytics.records(allocation_start, allocation_end)
        result = funnel_analytics().report(
            records,
            allocation_records,
            rates=rates,
            period_expenses=expenses,
            filters=filters,
            include_drafts=include_drafts,
        )
        result["period"] = {"from": start.isoformat(), "to": end.isoformat()}
        result["last_sync"] = analytics.last_sync()
        return jsonify(result)

    @app.get("/api/yclients/filters")
    @requires_auth
    def yclients_filters_api():
        return jsonify(yclients_analytics().filters())

    @app.get("/api/yclients")
    @requires_auth
    def yclients_api():
        start, end = _parse_yclients_period()
        result = yclients_analytics().report(
            start,
            end,
            filters=_filters(keys=("branch", "direction")),
            detail_limit=request.args.get("detail_limit", "100"),
        )
        result["last_sync"] = yclients_analytics().last_sync()
        return jsonify(result)

    @app.get("/api/yclients/plan-fact")
    @requires_auth
    def yclients_plan_fact_api():
        month = parse_plan_month(
            request.args.get("month") or _moscow_today().strftime("%Y-%m")
        )
        plans = plan_fact_repository().list_plans(month)
        result = plan_fact_analytics().report(
            month,
            plans,
            branches=tuple(
                value
                for value in request.args.getlist("branch")
                if value and value != "Все"
            ),
            directions=tuple(
                value
                for value in request.args.getlist("direction")
                if value and value != "Все"
            ),
        )
        result["options"] = {
            "branches": expense_branches(),
            "directions": list(PLAN_FACT_DIRECTIONS),
        }
        result["last_sync"] = yclients_analytics().last_sync()
        return jsonify(result)

    def yclients_color_result(export: bool = False):
        start, end = _parse_yclients_period()
        try:
            page = int(request.args.get("page", "1"))
            per_page = 100000 if export else int(request.args.get("per_page", "50"))
        except ValueError as error:
            raise ValueError("Некорректный номер страницы") from error
        return start, end, yclients_analytics().color_reconciliation(
            start,
            end,
            filters=_filters(keys=("branch", "direction")),
            visit_types=request.args.getlist("visit_type"),
            colors=request.args.getlist("record_color"),
            search=request.args.get("q", ""),
            page=page,
            per_page=per_page,
        )

    @app.get("/api/yclients/colors")
    @requires_auth
    def yclients_colors_api():
        start, end, result = yclients_color_result()
        result["period"] = {"from": start.isoformat(), "to": end.isoformat()}
        result["last_sync"] = yclients_analytics().last_sync()
        return jsonify(result)

    @app.get("/api/yclients/colors/export.csv")
    @requires_auth
    def yclients_colors_export_csv():
        start, end, result = yclients_color_result(export=True)
        payload = color_reconciliation_to_csv(result["rows"])
        return Response(
            payload,
            mimetype="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": (
                    f'attachment; filename="yclients_colors_{start}_{end}.csv"'
                )
            },
        )

    @app.post("/api/budgets")
    @requires_auth
    def budget_save_api():
        require_csrf("budget_csrf")
        saved = budget_repository().save_budget(
            json_payload(),
            branches=expense_branches(),
            user=str(session.get("dashboard_user", "dashboard")),
        )
        return jsonify({"budget": saved}), 201

    @app.put("/api/budgets/<int:budget_id>")
    @requires_auth
    def budget_update_api(budget_id: int):
        require_csrf("budget_csrf")
        saved = budget_repository().update_budget(
            budget_id,
            json_payload(),
            branches=expense_branches(),
            user=str(session.get("dashboard_user", "dashboard")),
        )
        return jsonify({"budget": saved})

    @app.delete("/api/budgets/<int:budget_id>")
    @requires_auth
    def budget_delete_api(budget_id: int):
        require_csrf("budget_csrf")
        deleted = budget_repository().delete_budget(
            budget_id,
            user=str(session.get("dashboard_user", "dashboard")),
        )
        return jsonify({"budget": deleted})

    @app.get("/api/dashboard")
    @requires_auth
    def dashboard():
        start, end = _parse_period()
        records = analytics.records(start, end)
        filters = _filters()
        result = analytics.aggregate(
            records,
            filters=filters,
        )
        result.pop("details", None)
        include_drafts = request.args.get("include_drafts", "false").lower() in {
            "1",
            "true",
            "yes",
        }
        try:
            repository = expense_repository()
            rates = repository.list_rates()
            expenses = repository.list_period_expenses(start, end)
            allocation_records = records
            if expenses:
                allocation_start = min(
                    date.fromisoformat(row["period_from"]) for row in expenses
                )
                allocation_end = max(
                    date.fromisoformat(row["period_to"]) for row in expenses
                )
                if allocation_start < start or allocation_end > end:
                    allocation_records = analytics.records(
                        min(start, allocation_start), max(end, allocation_end)
                    )
            result["marketing"] = calculate_marketing_performance(
                records,
                allocation_records,
                rates=rates,
                period_expenses=expenses,
                filters=filters,
                include_drafts=include_drafts,
            )
        except MarketingSchemaMissing:
            result["marketing"] = {
                "available": False,
                "kpi": {
                    "spend": None,
                    "leads": result["kpi"]["leads"],
                    "bookings": result["kpi"]["bookings"],
                    "conversion": result["kpi"]["conversion"],
                    "cpl": None,
                    "cost_per_booking": None,
                },
                "coverage": {
                    "status": "unavailable",
                    "label": "Раздел расходов ещё не настроен в базе",
                    "missing_units": [],
                    "unconfigured_sources": [],
                    "includes_drafts": include_drafts,
                },
                "groups": {
                    "source": [],
                    "branch": [],
                    "direction": [],
                    "source_branch": [],
                },
            }
        result["period"] = {
            "from": start.isoformat(),
            "to": end.isoformat(),
        }
        return jsonify(result)

    @app.get("/api/trends")
    @requires_auth
    def trends_api():
        start, end = _parse_period()
        granularity = request.args.get("granularity", "day")
        breakdown = request.args.get("breakdown", "none")
        compare = (
            request.args.get("compare", "false").lower() in {"1", "true", "yes"}
            and breakdown == "none"
        )
        include_drafts = request.args.get("include_drafts", "false").lower() in {
            "1",
            "true",
            "yes",
        }
        filters = _filters()
        current_records = analytics.records(start, end)
        previous_start = None
        previous_end = None
        previous_records = []
        if compare:
            duration = (end - start).days + 1
            candidate_start = start - timedelta(days=duration)
            if candidate_start >= MIN_DASHBOARD_DATE:
                previous_start = candidate_start
                previous_end = start - timedelta(days=1)
                previous_records = analytics.records(previous_start, previous_end)

        combined_start = previous_start or start
        try:
            repository = expense_repository()
            rates = repository.list_rates()
            expenses = repository.list_period_expenses(combined_start, end)
        except MarketingSchemaMissing:
            rates = []
            expenses = []

        allocation_start = combined_start
        allocation_end = end
        if expenses:
            allocation_start = min(
                allocation_start,
                min(date.fromisoformat(row["period_from"]) for row in expenses),
            )
            allocation_end = max(
                allocation_end,
                max(date.fromisoformat(row["period_to"]) for row in expenses),
            )
        allocation_records = analytics.records(allocation_start, allocation_end)
        current = calculate_trend_data(
            current_records,
            allocation_records,
            rates=rates,
            period_expenses=expenses,
            start=start,
            end=end,
            granularity=granularity,
            breakdown=breakdown,
            filters=filters,
            include_drafts=include_drafts,
        )
        previous = None
        if previous_start and previous_end:
            previous = calculate_trend_data(
                previous_records,
                allocation_records,
                rates=rates,
                period_expenses=expenses,
                start=previous_start,
                end=previous_end,
                granularity=granularity,
                breakdown="none",
                filters=filters,
                include_drafts=include_drafts,
            )
        return jsonify(
            {
                "current": current,
                "previous": previous,
                "comparison_available": previous is not None,
            }
        )

    def verification_result(*, export: bool = False):
        start, end = _parse_period()
        booking_start = _optional_date("booking_date_from")
        booking_end = _optional_date("booking_date_to")
        if booking_start and booking_end and booking_end < booking_start:
            raise ValueError("Дата записи «по» не может быть раньше даты «с»")
        try:
            page = int(request.args.get("page", "1"))
            per_page = 100000 if export else int(request.args.get("per_page", "50"))
        except ValueError as error:
            raise ValueError("Некорректный номер страницы") from error
        return analytics.verification(
            start,
            end,
            view=request.args.get("view", "leads"),
            filters=_filters(("branch", "source")),
            booking_start=booking_start,
            booking_end=booking_end,
            search=request.args.get("q", ""),
            sort=request.args.get("sort", "created_at"),
            order=request.args.get("order", "desc"),
            page=page,
            per_page=per_page,
        )

    @app.get("/api/verification")
    @requires_auth
    def verification_api():
        result = verification_result()
        result["period"] = {
            "from": request.args.get("date_from"),
            "to": request.args.get("date_to"),
        }
        return jsonify(result)

    @app.get("/api/verification/export.csv")
    @requires_auth
    def verification_export_csv():
        result = verification_result(export=True)
        view = result["view"]
        payload = verification_to_csv(result["rows"], view)
        return Response(
            payload,
            mimetype="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": (
                    f'attachment; filename="amo_{view}_verification.csv"'
                )
            },
        )

    @app.get("/api/export.csv")
    @requires_auth
    def export_csv():
        start, end = _parse_period()
        result = analytics.aggregate(
            analytics.records(start, end),
            filters=_filters(),
            search=request.args.get("q", ""),
        )
        payload = details_to_csv(result["details"], start, end)
        return Response(
            payload,
            mimetype="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": (
                    f'attachment; filename="amo_dashboard_{start}_{end}.csv"'
                )
            },
        )

    @app.get("/api/export.xlsx")
    @requires_auth
    def export_xlsx():
        start, end = _parse_period()
        result = analytics.aggregate(
            analytics.records(start, end),
            filters=_filters(),
            search=request.args.get("q", ""),
        )
        return send_file(
            _xlsx_bytes(result["details"], start, end),
            as_attachment=True,
            download_name=f"amo_dashboard_{start}_{end}.xlsx",
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    return app
